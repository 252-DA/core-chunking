"""
DocxParser — parse DOCX → ParsedDocument dùng python-docx.

Lợi thế so với PDF:
  - Headings được đánh dấu rõ ràng qua Word styles (Heading 1, Heading 2, ...)
  - Không cần heuristic font size → heading detection chính xác hơn nhiều
  - Tables được identify trực tiếp
  - Images extract từ relationships
"""
import io
import uuid
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.domain.entities.document import (
    Document,
    DocumentType,
    ElementType,
    ParsedDocument,
    Section,
)
from src.domain.ports.parser import IParser
from src.infrastructure.config import ParserConfig
from src.shared.logger import get_logger
from src.shared.result import Err, Ok, Result

logger = get_logger(__name__)

# Word style names → heading level
_HEADING_STYLE_MAP: dict[str, int] = {
    "heading 1": 1,
    "heading 2": 2,
    "heading 3": 3,
    "heading 4": 4,
    "heading 5": 5,
    "heading 6": 6,
    # Vietnamese Word thường đặt tên kiểu này
    "tiêu đề 1": 1,
    "tiêu đề 2": 2,
    "tiêu đề 3": 3,
}


class DocxParser(IParser):
    def __init__(self, config: ParserConfig) -> None:
        self._config = config

    @property
    def supported_types(self) -> tuple[DocumentType, ...]:
        return (DocumentType.DOCX,)

    def parse(self, path: Path) -> Result[ParsedDocument, Exception]:
        logger.info("docx_parser.started", file=path.name)
        try:
            doc = docx.Document(str(path))
            result = self._extract(path, doc)
            logger.info(
                "docx_parser.completed",
                file=path.name,
                sections=len(result.sections),
                images=len(result.images),
            )
            return Ok(result)
        except Exception as e:
            logger.error("docx_parser.failed", file=path.name, error=str(e))
            return Err(e)

    # ------------------------------------------------------------------
    # Core extraction
    # ------------------------------------------------------------------

    def _extract(self, path: Path, doc: docx.Document) -> ParsedDocument:
        sections: list[Section] = []
        images: dict[str, bytes] = {}

        # Lấy tất cả block-level elements (paragraphs + tables) theo thứ tự
        body_elements = self._iter_body_elements(doc)

        for element in body_elements:
            if isinstance(element, Paragraph):
                section = self._parse_paragraph(element)
                if section:
                    sections.append(section)
            elif isinstance(element, Table):
                section = self._parse_table(element)
                if section:
                    sections.append(section)

        # Extract images từ relationships
        images = self._extract_images(doc)

        document = Document(
            id=str(uuid.uuid4()),
            name=path.name,
            path=path,
            doc_type=DocumentType.DOCX,
            size_bytes=path.stat().st_size,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        return ParsedDocument(
            document=document,
            sections=sections,
            page_count=0,  # DOCX không có page count trực tiếp — để 0
            images=images,
        )

    # ------------------------------------------------------------------
    # Element iteration
    # ------------------------------------------------------------------

    def _iter_body_elements(self, doc: docx.Document):
        """
        Yield paragraphs và tables theo đúng thứ tự xuất hiện trong document.
        python-docx có doc.paragraphs và doc.tables riêng nhưng không giữ thứ tự.
        Cần iterate qua XML body để giữ đúng thứ tự.
        """
        body = doc.element.body
        for child in body.iterchildren():
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                yield Paragraph(child, doc)
            elif tag == "tbl":
                yield Table(child, doc)

    # ------------------------------------------------------------------
    # Paragraph parsing
    # ------------------------------------------------------------------

    def _parse_paragraph(self, para: Paragraph) -> Section | None:
        text = para.text.strip()
        if not text:
            return None

        heading_level = self._detect_heading_level(para)

        if heading_level > 0:
            return Section(
                content=text,
                element_type=ElementType.HEADING,
                heading=text,
                heading_level=heading_level,
            )

        return Section(
            content=text,
            element_type=ElementType.PARAGRAPH,
        )

    def _detect_heading_level(self, para: Paragraph) -> int:
        """
        Detect heading level từ Word style name.
        Chính xác hơn nhiều so với heuristic font size.
        """
        if not para.style:
            return 0

        style_name = para.style.name.lower().strip()

        # Check map trực tiếp
        if style_name in _HEADING_STYLE_MAP:
            return _HEADING_STYLE_MAP[style_name]

        # Check dạng "heading X" generic
        if style_name.startswith("heading "):
            try:
                level = int(style_name.split(" ")[-1])
                return min(level, 6)
            except ValueError:
                pass

        # Check outline level từ XML (paragraph properties)
        try:
            pPr = para._element.pPr
            if pPr is not None:
                outlineLvl = pPr.find(qn("w:outlineLvl"))
                if outlineLvl is not None:
                    val = int(outlineLvl.get(qn("w:val"), 9))
                    if val < 9:
                        return val + 1  # 0-based → 1-based
        except Exception:
            pass

        return 0

    # ------------------------------------------------------------------
    # Table parsing
    # ------------------------------------------------------------------

    def _parse_table(self, table: Table) -> Section | None:
        """Convert table thành plain text dạng markdown-ish."""
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(cells)

        if not rows:
            return None

        # Header row + separator + data rows
        lines = []
        if rows:
            lines.append(" | ".join(rows[0]))
            lines.append(" | ".join(["---"] * len(rows[0])))
        for row in rows[1:]:
            lines.append(" | ".join(row))

        return Section(
            content="\n".join(lines),
            element_type=ElementType.TABLE,
        )

    # ------------------------------------------------------------------
    # Image extraction
    # ------------------------------------------------------------------

    def _extract_images(self, doc: docx.Document) -> dict[str, bytes]:
        images: dict[str, bytes] = {}
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                try:
                    img_part = rel.target_part
                    ext = img_part.content_type.split("/")[-1]
                    filename = f"img_{rel.rId}.{ext}"
                    images[filename] = img_part.blob
                except Exception:
                    pass
        return images
