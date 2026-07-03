"""
PptxParser — parse PPTX → ParsedDocument dùng python-pptx.

Strategy:
  - Mỗi slide → 1 Section (element_type=SLIDE)
  - Title của slide → heading
  - Body text + notes → content
  - Images extract từ shapes trong mỗi slide

Khác với PDF/DOCX, PPTX có natural boundary là slide —
không cần detect heading, mỗi slide là 1 đơn vị content.
"""
import io
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER_TYPE
from pptx.util import Pt

from document_chunk.domain.entities.document import (
    Document,
    DocumentType,
    ElementType,
    ParsedDocument,
    Section,
)
from document_chunk.domain.exceptions import ParseError
from document_chunk.domain.ports.parser import IParser
from document_chunk.infrastructure.config import ParserConfig
from document_chunk.shared.logger import get_logger
from document_chunk.shared.result import Err, Ok, Result
from document_chunk.shared.tracing import get_tracer

logger = get_logger(__name__)
tracer = get_tracer(__name__)


class PptxParser(IParser):
    def __init__(self, config: ParserConfig) -> None:
        self._config = config

    @property
    def supported_types(self) -> tuple[DocumentType, ...]:
        return (DocumentType.PPTX,)

    def parse(self, path: Path) -> Result[ParsedDocument, Exception]:
        logger.info("pptx_parser.started", file=path.name)
        with tracer.start_as_current_span("pptx_parser.parse") as span:
            span.set_attribute("file.name", path.name)
            span.set_attribute("file.size_bytes", path.stat().st_size)
            try:
                prs = Presentation(str(path))
                result = self._extract(path, prs)
                span.set_attribute("slides", result.page_count)
                span.set_attribute("sections", len(result.sections))
                span.set_attribute("images", len(result.images))
                logger.info(
                    "pptx_parser.completed",
                    file=path.name,
                    slides=result.page_count,
                    sections=len(result.sections),
                    images=len(result.images),
                )
                return Ok(result)
            except Exception as e:
                logger.error("pptx_parser.failed", file=path.name, error=str(e))
                return Err(ParseError(f"Failed to parse PPTX '{path.name}'", cause=e))

    # ------------------------------------------------------------------
    # Core extraction
    # ------------------------------------------------------------------

    def _extract(self, path: Path, prs: Presentation) -> ParsedDocument:
        sections: list[Section] = []
        images: dict[str, bytes] = {}
        slide_count = len(prs.slides)

        for slide_num, slide in enumerate(prs.slides, start=1):
            title = self._get_slide_title(slide)
            body_text = self._get_body_text(slide)
            notes_text = self._get_notes_text(slide)
            slide_images = self._extract_slide_images(slide, slide_num, images)

            # Build content: body + notes (nếu có)
            content_parts = [p for p in [body_text, notes_text] if p]
            content = "\n\n".join(content_parts)

            # Bỏ qua slide trống hoàn toàn
            if not title and not content:
                continue

            sections.append(Section(
                content=content or title or "",
                element_type=ElementType.SLIDE,
                heading=title,
                heading_level=1 if title else 0,
                page_number=slide_num,
                images=tuple(slide_images),
                metadata={"slide_number": slide_num, "total_slides": slide_count},
            ))

        document = Document(
            id=str(uuid.uuid4()),
            name=path.name,
            path=path,
            doc_type=DocumentType.PPTX,
            size_bytes=path.stat().st_size,
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

        return ParsedDocument(
            document=document,
            sections=sections,
            page_count=slide_count,
            images=images,
        )

    # ------------------------------------------------------------------
    # Slide content extraction
    # ------------------------------------------------------------------

    def _get_slide_title(self, slide) -> str | None:
        """Lấy title từ title placeholder hoặc shape đầu tiên có text."""
        # Ưu tiên title placeholder
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if shape.is_placeholder:
                ph_type = shape.placeholder_format.type
                if ph_type in (
                    PP_PLACEHOLDER_TYPE.TITLE,
                    PP_PLACEHOLDER_TYPE.CENTER_TITLE,
                ):
                    text = shape.text_frame.text.strip()
                    if text:
                        return text

        # Fallback: shape đầu tiên có text
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    return text

        return None

    def _get_body_text(self, slide) -> str:
        """
        Lấy toàn bộ text từ body/content placeholders.
        Bỏ qua title placeholder (đã lấy riêng).
        """
        texts: list[str] = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if shape.is_placeholder:
                ph_type = shape.placeholder_format.type
                # Bỏ qua title placeholder
                if ph_type in (
                    PP_PLACEHOLDER_TYPE.TITLE,
                    PP_PLACEHOLDER_TYPE.CENTER_TITLE,
                ):
                    continue

            # Extract từng paragraph trong text frame
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    texts.append(text)

        return "\n".join(texts)

    def _get_notes_text(self, slide) -> str | None:
        """Extract notes từ notes slide nếu có."""
        if not self._config.pptx_include_notes:
            return None

        try:
            notes_slide = slide.notes_slide
            if notes_slide and notes_slide.notes_text_frame:
                text = notes_slide.notes_text_frame.text.strip()
                if text:
                    return f"[Notes] {text}"
        except Exception:
            pass

        return None

    # ------------------------------------------------------------------
    # Image extraction
    # ------------------------------------------------------------------

    def _extract_slide_images(
        self,
        slide,
        slide_num: int,
        images: dict[str, bytes],
    ) -> list[str]:
        """Extract images từ shapes trong slide. Trả về list filenames."""
        slide_image_names: list[str] = []

        for shape in slide.shapes:
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE = 13
                try:
                    img = shape.image
                    ext = img.ext
                    filename = f"slide_{slide_num}_{shape.shape_id}.{ext}"
                    images[filename] = img.blob
                    slide_image_names.append(filename)
                except Exception as exc:
                    logger.warning(
                        "pptx_parser.image_extract_failed",
                        slide_number=slide_num,
                        shape_id=shape.shape_id,
                        error=str(exc),
                    )

        return slide_image_names
