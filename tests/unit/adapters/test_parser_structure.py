import fitz
import docx
from document_chunk.adapters.parsers.pdf_parser import PdfParser
from document_chunk.adapters.parsers.markdown_parser import MarkdownParser
from document_chunk.adapters.parsers.docx_parser import DocxParser
from document_chunk.domain.entities.document import ElementType
from document_chunk.infrastructure.config import ParserConfig


def test_pdf_short_styled_span_and_paragraphs(tmp_path):
    path = tmp_path / "styled.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page()
        x = 72
        for text, font in [
            ("He quan tri ", "helv"),
            ("CSDL", "hebo"),
            (" la phan mem quan ly du lieu.", "helv"),
        ]:
            page.insert_text((x, 72), text, fontname=font, fontsize=12)
            x += fitz.get_text_length(text, fontname=font, fontsize=12)
        page.insert_text((72, 120), "Doan van thu hai can duoc giu rieng.", fontsize=12)
        pdf.save(path)
    parsed = PdfParser(ParserConfig()).parse(path).unwrap()
    assert "CSDL" in parsed.sections[0].content
    assert len(parsed.sections) == 2
    assert parsed.metadata["parser"] == "pymupdf"


def test_pdf_running_headers(tmp_path):
    path = tmp_path / "headers.pdf"
    with fitz.open() as pdf:
        for i in range(4):
            page = pdf.new_page()
            page.insert_text((72, 25), f"Running header {i}")
            page.insert_text((72, 120), f"Body text on page {i}")
        pdf.save(path)
    parsed = PdfParser(ParserConfig()).parse(path).unwrap()
    assert len(parsed.sections) == 4
    assert all("Running" not in s.content for s in parsed.sections)


def test_markdown_fences_table_list_frontmatter(tmp_path):
    path = tmp_path / "test.md"
    path.write_text(
        "---\ntitle: ignore\n---\n# Real\n```python\n# comment\n    print(1)\n```\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n\n- One\n- Two\n"
    )
    parsed = MarkdownParser(ParserConfig()).parse(path).unwrap()
    assert [s.element_type for s in parsed.sections] == [
        ElementType.HEADING,
        ElementType.CODE,
        ElementType.TABLE,
        ElementType.LIST,
    ]
    assert "    print(1)" in parsed.sections[1].content


def test_markdown_tilde_fence_and_unclosed(tmp_path):
    path = tmp_path / "test.md"
    path.write_text("~~~~\n# comment\n~~~\n")
    parsed = MarkdownParser(ParserConfig()).parse(path).unwrap()
    assert len(parsed.sections) == 1
    assert parsed.sections[0].element_type == ElementType.CODE


def test_docx_list_and_merged_cells(tmp_path):
    path = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_paragraph("Item", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    doc.save(path)
    parsed = DocxParser(ParserConfig()).parse(path).unwrap()
    assert parsed.sections[0].element_type == ElementType.LIST
    assert parsed.sections[1].content.count("Merged") == 1


def test_pptx_preserves_table_blocks(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from document_chunk.adapters.parsers.pptx_parser import PptxParser
    prs = Presentation()
    prs.core_properties.title = 'Lesson SQL'
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = 'Examples'
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    table.cell(0, 0).text = 'Name'
    table.cell(0, 1).text = 'Value'
    table.cell(1, 0).text = 'SQL'
    table.cell(1, 1).text = 'Database'
    path = tmp_path/'slides.pptx'
    prs.save(path)
    parsed = PptxParser(ParserConfig()).parse(path).unwrap()
    assert parsed.metadata['title'] == 'Lesson SQL'
    assert parsed.sections[0].metadata['blocks'][0]['kind'] == 'table'
    assert 'Database' in parsed.sections[0].content
